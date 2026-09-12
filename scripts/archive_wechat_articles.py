from __future__ import annotations

import argparse
import io
import json
import re
import shutil
import sys
import urllib.parse
import urllib.request
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import html
from PIL import Image, ImageFile

ImageFile.LOAD_TRUNCATED_IMAGES = True

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/140.0.0.0 Safari/537.36",
    "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.7",
    "Referer": "https://mp.weixin.qq.com/",
}
SKIP_TEXT = {
    "视频", "小程序", "赞", "轻点两下取消赞", "在看", "轻点两下取消在看",
    "分享", "留言", "收藏", "听过", "写留言", "向上滑动看下一个",
}


def clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\u00a0", " ").replace("\u200b", "")).strip()


def safe_name(value: str) -> str:
    value = re.sub(r'[<>:"/\\|?*]', "_", value).strip().rstrip(".")
    return value or "未命名"


def normalize_wechat_url(url: str) -> str:
    parsed = urllib.parse.urlsplit(url)
    if parsed.netloc.lower() != "mp.weixin.qq.com":
        raise ValueError(f"not a WeChat article URL: {url}")
    pairs = urllib.parse.parse_qsl(parsed.query, keep_blank_values=True)
    keys = {key for key, _ in pairs}
    if "biz" in keys and "__biz" not in keys:
        pairs = [("__biz" if key == "biz" else key, value) for key, value in pairs]
    query = urllib.parse.urlencode(pairs, doseq=True)
    return urllib.parse.urlunsplit(("https", parsed.netloc, parsed.path, query, ""))


def fetch(url: str, referer: str | None = None) -> tuple[bytes, str]:
    headers = dict(HEADERS)
    if referer:
        headers["Referer"] = referer
    request = urllib.request.Request(url, headers=headers)
    with urllib.request.urlopen(request, timeout=90) as response:
        return response.read(), response.geturl()


def decode_html(data: bytes) -> str:
    for encoding in ("utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def validate_tree(tree) -> tuple[object, str, str, str]:
    roots = tree.xpath("//*[@id='js_content']")
    title = clean_text(tree.xpath("string(//*[@id='activity-name'])"))
    if not roots or not title:
        page_text = clean_text(tree.xpath("string(//body)"))[:200]
        raise ValueError(f"invalid or blocked WeChat article page: {page_text}")
    body_text = clean_text(" ".join(roots[0].itertext()))
    image_count = len(roots[0].xpath(".//img"))
    if not body_text and image_count == 0:
        raise ValueError("WeChat article body is empty")
    author = clean_text(tree.xpath("string(//*[@id='js_name'])"))
    date = clean_text(tree.xpath("string(//*[@id='publish_time'])"))
    return roots[0], title, author, date


def image_url(element, base_url: str) -> str | None:
    for key in ("data-src", "data-original", "data-echo", "src"):
        value = element.get(key)
        if value and not value.startswith(("data:", "blob:")):
            return urllib.parse.urljoin(base_url, "https:" + value if value.startswith("//") else value)
    return None


def extract_blocks(root, base_url: str) -> list[tuple[str, str]]:
    blocks: list[tuple[str, str]] = []
    seen_text: set[str] = set()
    for element in root.iter():
        tag = element.tag.lower() if isinstance(element.tag, str) else ""
        if tag in {"script", "style", "noscript", "svg", "button"}:
            continue
        if tag == "img":
            url = image_url(element, base_url)
            if url:
                blocks.append(("image", url))
            continue
        if tag in {"p", "h1", "h2", "h3", "h4", "li"}:
            text = clean_text(" ".join(element.itertext()))
            if not text or text in SKIP_TEXT or text in seen_text or len(text) > 2500:
                continue
            seen_text.add(text)
            blocks.append(("heading" if tag.startswith("h") else "text", text))
        elif tag in {"section", "div"} and not element.xpath(".//p|.//h1|.//h2|.//h3|.//h4|.//li"):
            text = clean_text(" ".join(element.itertext()))
            if 1 < len(text) <= 300 and text not in SKIP_TEXT and text not in seen_text:
                seen_text.add(text)
                blocks.append(("text", text))
    compact: list[tuple[str, str]] = []
    for block in blocks:
        if not compact or block != compact[-1]:
            compact.append(block)
    return compact


def download_image(url: str, folder: Path, index: int, referer: str) -> Path | None:
    try:
        lower = url.lower()
        if "mmbiz_gif" in lower or "wx_fmt=gif" in lower:
            return None
        data, _ = fetch(url, referer)
        if len(data) < 700:
            return None
        with Image.open(io.BytesIO(data)) as source:
            source.seek(0)
            source.load()
            if source.width < 80 or source.height < 60:
                return None
            if source.mode not in ("RGB", "RGBA"):
                source = source.convert("RGBA" if "transparency" in source.info else "RGB")
            extension = ".png" if source.mode == "RGBA" else ".jpg"
            path = folder / f"图片{index:02d}{extension}"
            if extension == ".jpg":
                source.convert("RGB").save(path, "JPEG", quality=92, optimize=True)
            else:
                source.save(path, "PNG", optimize=True)
        return path
    except Exception as exc:
        print(f"WARN image {index}: {url[:100]}: {exc}", file=sys.stderr)
        return None


def set_font(run, size: float, bold: bool | None = None, color: tuple[int, int, int] | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    fonts = run._element.get_or_add_rPr().rFonts
    for key in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        fonts.set(qn(key), "Microsoft YaHei")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    rid = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "4F6B8A")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def new_document() -> Document:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.72)
    section.left_margin = section.right_margin = Inches(0.82)
    for style_name, size in (("Normal", 11), ("Title", 22), ("Heading 1", 15)):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Normal"].paragraph_format.space_after = Pt(7)
    document.styles["Normal"].paragraph_format.line_spacing = 1.45
    return document


def add_front(document: Document, label: str, title: str, author: str, date: str, url: str) -> None:
    paragraph = document.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(paragraph.add_run(label), 22, True)
    details = [value for value in (title if title != label else "", author, date) if value]
    if details:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(paragraph.add_run("  |  ".join(details)), 9.5, color=(90, 90, 90))
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(14)
    add_hyperlink(paragraph, "原始网页", url)


def add_text(document: Document, text: str, heading: bool) -> None:
    paragraph = document.add_paragraph(style="Heading 1" if heading else None)
    if heading:
        paragraph.paragraph_format.keep_with_next = True
        set_font(paragraph.add_run(text), 15, True)
    else:
        paragraph.paragraph_format.first_line_indent = Inches(0.3)
        paragraph.paragraph_format.widow_control = True
        set_font(paragraph.add_run(text), 11)


def add_image(document: Document, path: Path) -> None:
    with Image.open(path) as image:
        width_px, height_px = image.size
    width = min(6.6, 8.2 * width_px / max(height_px, 1))
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_together = True
    paragraph.add_run().add_picture(str(path), width=Inches(width))


def check_target(output_parent: Path, archive_name: str) -> Path:
    if not archive_name.strip() or Path(archive_name).name != archive_name or archive_name in {".", ".."}:
        raise ValueError("archive_name must be a safe leaf name, not a path")
    return output_parent.resolve() / safe_name(archive_name)


def build(args: argparse.Namespace) -> Path:
    payload = json.loads(Path(args.input).resolve().read_text(encoding="utf-8-sig"))
    articles = payload.get("articles")
    if not isinstance(articles, list) or not articles:
        raise ValueError("articles must be a non-empty array")
    output_parent = Path(args.output_parent).resolve()
    output_parent.mkdir(parents=True, exist_ok=True)
    archive = check_target(output_parent, str(payload.get("archive_name", "")))
    cache = Path(args.cache_dir).resolve()
    cache.mkdir(parents=True, exist_ok=True)
    if archive.exists():
        if not args.overwrite:
            raise FileExistsError(f"archive already exists: {archive}; use --overwrite intentionally")
        if archive == output_parent or output_parent not in archive.parents:
            raise ValueError("refusing to replace an unsafe archive path")
        shutil.rmtree(archive)
    archive.mkdir()
    seen_ids: set[str] = set()
    records: list[dict[str, object]] = []
    for position, article in enumerate(articles, 1):
        item_id = str(article.get("id", "")).strip()
        label = str(article.get("label", "")).strip()
        original_url = str(article.get("url", "")).strip()
        if not item_id or not label or not original_url:
            raise ValueError(f"article {position} requires id, label, and url")
        if item_id in seen_ids:
            raise ValueError(f"duplicate article id: {item_id}")
        seen_ids.add(item_id)
        fetch_url = normalize_wechat_url(original_url)
        if article.get("html_path"):
            data = Path(str(article["html_path"])).resolve().read_bytes()
            final_url = fetch_url
        else:
            data, final_url = fetch(fetch_url)
        (cache / f"{safe_name(item_id)}.html").write_bytes(data)
        tree = html.fromstring(decode_html(data))
        root, source_title, author, date = validate_tree(tree)
        blocks = extract_blocks(root, final_url)
        folder_name = f"{safe_name(item_id)}-{safe_name(label)}"
        folder = archive / folder_name
        images_folder = folder / "图片"
        images_folder.mkdir(parents=True)
        document = new_document()
        add_front(document, label, source_title, author, date, original_url)
        image_attempt = archived_images = text_blocks = skipped_dynamic = 0
        for kind, value in blocks:
            if kind == "image":
                image_attempt += 1
                path = download_image(value, images_folder, image_attempt, final_url)
                if path:
                    add_image(document, path)
                    archived_images += 1
                elif "gif" in value.lower():
                    skipped_dynamic += 1
            elif value != source_title:
                add_text(document, value, kind == "heading")
                text_blocks += 1
        docx_path = folder / f"{folder_name}.docx"
        document.save(docx_path)
        checked = Document(docx_path)
        image_files = [path for path in images_folder.iterdir() if path.is_file()]
        embedded = len(checked.inline_shapes)
        if not checked.paragraphs or embedded != len(image_files) or embedded != archived_images:
            raise RuntimeError(f"structural validation failed for {item_id}")
        records.append({
            "编号": item_id,
            "标签": label,
            "源标题": source_title,
            "原始链接": original_url,
            "最终读取链接": final_url,
            "文件夹": folder_name,
            "Word文件": docx_path.name,
            "文本块数": text_blocks,
            "图片文件数": len(image_files),
            "Word内嵌图片数": embedded,
            "跳过动态图片数": skipped_dynamic,
            "结构校验": "通过",
        })
        print(f"{item_id}\ttext={text_blocks}\timages={embedded}\t{docx_path}")
    manifest = {
        "归档名称": payload["archive_name"],
        "文章数量": len(records),
        "总图片数": sum(int(record["图片文件数"]) for record in records),
        "项目": records,
    }
    (archive / "归档清单.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    if args.zip:
        zip_path = output_parent / f"{archive.name}.zip"
        if zip_path.exists():
            if not args.overwrite:
                raise FileExistsError(f"ZIP already exists: {zip_path}; use --overwrite intentionally")
            zip_path.unlink()
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as bundle:
            for path in sorted(archive.rglob("*")):
                if path.is_file():
                    bundle.write(path, path.relative_to(output_parent))
        with zipfile.ZipFile(zip_path) as bundle:
            if bundle.testzip() is not None:
                raise RuntimeError("ZIP integrity check failed")
        print(f"ZIP\t{zip_path}")
    return archive


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Archive WeChat articles as illustrated Word documents.")
    parser.add_argument("--input", required=True, help="UTF-8 JSON input file")
    parser.add_argument("--output-parent", required=True, help="Parent directory for archive folder and ZIP")
    parser.add_argument("--cache-dir", required=True, help="Temporary directory for fetched HTML")
    parser.add_argument("--zip", action="store_true", help="Create a ZIP beside the archive folder")
    parser.add_argument("--overwrite", action="store_true", help="Replace only the exact named archive directory")
    return parser.parse_args()


if __name__ == "__main__":
    try:
        build(parse_args())
    except Exception as error:
        print(f"ERROR: {error}", file=sys.stderr)
        raise SystemExit(1)
